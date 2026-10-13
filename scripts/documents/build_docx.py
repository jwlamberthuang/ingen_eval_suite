"""Build an editable methodology document from a local JSON document source.

Run with the bundled document runtime. Source writeups and output documents
can live outside the software checkout; only authoring code lives here.
"""

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from set_bilingual_fonts import apply_bilingual_fonts


FONT_NAME = "Times New Roman"


def explicit_font(props, font_name=FONT_NAME):
    """Set all Word font slots and remove theme references that override them."""
    fonts = props.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.insert(0, fonts)
    fonts.attrib.clear()
    for slot in ["ascii", "hAnsi", "eastAsia", "cs"]:
        fonts.set(qn("w:" + slot), font_name)


def set_font(style, size, bold=False):
    style.font.name = FONT_NAME
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    props = style._element.get_or_add_rPr()
    explicit_font(props)
    size_cs = props.find(qn("w:szCs"))
    if size_cs is None:
        size_cs = OxmlElement("w:szCs")
        props.append(size_cs)
    size_cs.set(qn("w:val"), str(int(size * 2)))


def inline(paragraph, text):
    """Support bold text, inline code, and external Markdown links."""
    pattern = r"(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\(https?://[^)]+\))"
    for token in re.split(pattern, text):
        if not token:
            continue
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("["):
            match = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
            if not match:
                paragraph.add_run(token)
                continue
            label, url = match.groups()
            link = OxmlElement("w:hyperlink")
            rel = paragraph.part.relate_to(
                url,
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                is_external=True,
            )
            link.set(qn("r:id"), rel)
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "17365D")
            props.append(color)
            underline = OxmlElement("w:u")
            underline.set(qn("w:val"), "single")
            props.append(underline)
            run.append(props)
            value = OxmlElement("w:t")
            value.text = label
            run.append(value)
            link.append(run)
            paragraph._p.append(link)
        else:
            paragraph.add_run(token)


def math_node(value):
    """Construct simple native Word mathematics without rasterizing it."""
    if isinstance(value, str):
        run = OxmlElement("m:r")
        text = OxmlElement("m:t")
        text.text = value
        run.append(text)
        return run
    if "sum" in value:
        spec = value["sum"]
        node = OxmlElement("m:nary")
        props = OxmlElement("m:naryPr")
        symbol = OxmlElement("m:chr")
        symbol.set(qn("m:val"), "∑")
        props.append(symbol)
        location = OxmlElement("m:limLoc")
        location.set(qn("m:val"), "subSup")
        props.append(location)
        hidden = OxmlElement("m:supHide")
        hidden.set(qn("m:val"), "1")
        props.append(hidden)
        node.append(props)
        sub = OxmlElement("m:sub")
        sub.append(math_node(spec["subscript"]))
        node.append(sub)
        sup = OxmlElement("m:sup")
        node.append(sup)
        body = OxmlElement("m:e")
        append_math(body, spec["body"])
        node.append(body)
        return node
    if "sub" in value:
        spec = value["sub"]
        node = OxmlElement("m:sSub")
        base = OxmlElement("m:e")
        base.append(math_node(spec["base"]))
        sub = OxmlElement("m:sub")
        sub.append(math_node(spec["index"]))
        node.extend([base, sub])
        return node
    raise ValueError(f"Unsupported equation node: {value}")


def append_math(parent, value):
    if isinstance(value, dict) and "sequence" in value:
        for child in value["sequence"]:
            append_math(parent, child)
    else:
        parent.append(math_node(value))


def add_equation(doc, spec):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.keep_with_next = True
    container = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    math.append(math_node(spec["lhs"] + " = "))
    fraction = OxmlElement("m:f")
    numerator = OxmlElement("m:num")
    append_math(numerator, spec["numerator"])
    denominator = OxmlElement("m:den")
    append_math(denominator, spec["denominator"])
    fraction.extend([numerator, denominator])
    math.append(fraction)
    container.append(math)
    paragraph._p.append(container)


def add_table(doc, block):
    if block.get("caption"):
        paragraph = doc.add_paragraph(style="Caption")
        inline(paragraph, block["caption"])
        paragraph.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=len(block["headers"]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = block.get("widths", [6.9 / len(block["headers"])] * len(block["headers"]))
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement("w:" + side)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    props.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side, value in [("top", 90), ("bottom", 90), ("left", 110), ("right", 110)]:
        margin = OxmlElement("w:" + side)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    props.append(margins)
    rows = [block["headers"], *block["rows"]]
    for index, row_values in enumerate(rows):
        row = table.rows[0] if index == 0 else table.add_row()
        row_props = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        row_props.append(no_split)
        if index == 0:
            repeat = OxmlElement("w:tblHeader")
            row_props.append(repeat)
        for col, (cell, value) in enumerate(zip(row.cells, row_values)):
            cell.width = Inches(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E5EBF2" if index == 0 else ("F6F8FA" if index % 2 == 0 else "FFFFFF"))
            cell._tc.get_or_add_tcPr().append(shading)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.06
            if block.get("keep_together") and index < len(rows) - 1:
                paragraph.paragraph_format.keep_with_next = True
            if col in block.get("center_columns", []):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(paragraph, str(value))
            for run in paragraph.runs:
                run.font.size = Pt(block.get("font_size", 11))
                if index == 0:
                    run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)


def normalize_fonts(doc):
    """Make inherited, linked, direct, and theme fonts agree in Microsoft Word."""
    styles_root = doc.styles._element
    for style in doc.styles:
        explicit_font(style._element.get_or_add_rPr())
    defaults = styles_root.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        styles_root.insert(0, defaults)
    run_default = defaults.find(qn("w:rPrDefault"))
    if run_default is None:
        run_default = OxmlElement("w:rPrDefault")
        defaults.append(run_default)
    props = run_default.find(qn("w:rPr"))
    if props is None:
        props = OxmlElement("w:rPr")
        run_default.append(props)
    explicit_font(props)
    for fonts in styles_root.iter(qn("w:rFonts")):
        explicit_font(fonts.getparent())
    roots = [doc._element, doc.part.numbering_part.element]
    for section in doc.sections:
        roots.extend([section.header._element, section.footer._element])
    for root in roots:
        for run in root.iter(qn("w:r")):
            props = run.find(qn("w:rPr"))
            if props is None:
                props = OxmlElement("w:rPr")
                run.insert(0, props)
            explicit_font(props)
        for fonts in root.iter(qn("w:rFonts")):
            explicit_font(fonts.getparent())
    # Word's default bullets use a Symbol private-use character. Replace it
    # with an ordinary Unicode bullet before assigning the requested text font.
    for level in doc.part.numbering_part.element.iter(qn("w:lvl")):
        fmt = level.find(qn("w:numFmt"))
        if fmt is not None and fmt.get(qn("w:val")) == "bullet":
            label = level.find(qn("w:lvlText"))
            if label is not None:
                label.set(qn("w:val"), "•")
    from lxml import etree
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    for part in doc.part.package.parts:
        if part.content_type == "application/vnd.openxmlformats-officedocument.theme+xml":
            theme = etree.fromstring(part.blob)
            for group in ["majorFont", "minorFont"]:
                for slot in ["latin", "ea", "cs"]:
                    item = theme.find(f".//{{{a}}}{group}/{{{a}}}{slot}")
                    if item is not None:
                        item.set("typeface", FONT_NAME)
            part._blob = etree.tostring(theme, xml_declaration=True, encoding="UTF-8", standalone=True)


def build(source, output):
    data = json.loads(source.read_text())
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.footer_distance = Inches(0.3)
    styles = doc.styles
    set_font(styles["Normal"], 12)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Normal"].paragraph_format.line_spacing = 1.1
    for name, size, after in [("Title", 24, 10), ("Subtitle", 14, 10), ("Heading 1", 16, 8), ("Heading 2", 13, 6), ("Heading 3", 12, 5)]:
        set_font(styles[name], size, name != "Subtitle")
        styles[name].paragraph_format.space_before = Pt(12 if name.startswith("Heading") else 0)
        styles[name].paragraph_format.space_after = Pt(after)
        styles[name].paragraph_format.keep_with_next = True
        props = styles[name]._element.get_or_add_pPr()
        for border in list(props.findall(qn("w:pBdr"))):
            props.remove(border)
    set_font(styles["Caption"], 11)
    styles["Caption"].font.italic = False
    styles["Caption"].paragraph_format.space_before = Pt(3)
    styles["Caption"].paragraph_format.space_after = Pt(5)
    for name in ["List Bullet", "List Number"]:
        set_font(styles[name], 12)
        styles[name].paragraph_format.space_after = Pt(4)
        styles[name].paragraph_format.line_spacing = 1.1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Page ").font.size = Pt(10)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = data["title"]
    doc.core_properties.subject = data.get("subject", "Evaluation methodology")
    doc.core_properties.author = data.get("author", "")
    doc.core_properties.version = data.get("version", "0.1")
    for block in data["blocks"]:
        kind = block["type"]
        if kind in ["title", "subtitle"]:
            paragraph = doc.add_paragraph(style=kind.capitalize())
            inline(paragraph, block["text"])
        elif kind == "heading":
            paragraph = doc.add_heading(level=block.get("level", 1))
            inline(paragraph, block["text"])
        elif kind == "paragraph":
            paragraph = doc.add_paragraph()
            inline(paragraph, block["text"])
        elif kind == "list":
            for item in block["items"]:
                paragraph = doc.add_paragraph(style="List Number" if block.get("ordered") else "List Bullet")
                inline(paragraph, item)
        elif kind == "table":
            add_table(doc, block)
        elif kind == "equation":
            add_equation(doc, block)
        elif kind == "page_break":
            doc.add_page_break()
        else:
            raise ValueError(f"Unsupported block type: {kind}")
    output.parent.mkdir(parents=True, exist_ok=True)
    normalize_fonts(doc)
    doc.save(output)
    apply_bilingual_fonts(output)
    print(f"Saved {output}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source, args.output)
